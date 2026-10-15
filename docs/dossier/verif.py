# -*- coding: utf-8 -*-
"""Contrôle de complétude : tout le contenu rédigé se retrouve-t-il dans le
dossier rempli ?

Les écarts attendus sont déclarés ici et comptés à part :
 - les titres repris par la trame, qui ne sont pas recopiés ;
 - les blocs retirés pour tenir dans les 60 pages (code en double, etc.) ;
 - les dessins ASCII remplacés par de vraies figures ;
 - les renvois « § 5.2.2 » ramenés au niveau le plus profond que la trame
   connaisse (« § 5.2 ») : le texte change, le paragraphe est bien là.
"""
import re, unicodedata, sys
import docx
from docx.oxml.ns import qn

SRC = "/home/user/HelpMeDraft---projet-fil-rouge/Dossier projet CDA - HelpMeDraft.docx"
OUT = "/home/user/HelpMeDraft---projet-fil-rouge/docs/dossier/dossier-projet-rempli.docx"

def cle(t, sans_renvoi=False):
    if sans_renvoi:
        t = re.sub(r"§\s*[\d\.]+", "§", t)
    t = unicodedata.normalize("NFKD", t)
    t = "".join(c for c in t if not unicodedata.combining(c))
    return re.sub(r"[^a-z0-9]+", "", t.lower())

src, out = docx.Document(SRC), docx.Document(OUT)

TITRES_TRAME = {27, 51, 60, 82, 86, 102, 121, 159, 199, 221, 236, 256, 300,
                331, 360, 394, 539, 581, 582, 604, 609, 614, 623}
COUPES = ({334, 341, 347, 352, 363}                   # code déjà en annexe
          | {509}                                     # paragraphe en double du § 9
          | set(range(562, 573))                        # veille AI Act condensée
          | set(range(282, 287))                        # lancement en développement
          | {147}                                       # wireframe ASCII de l'éditeur
          | {299, 419, 467, 516})                       # intitulés déjà portés par la trame
ASCII_REMPLACES = {105, 139, 164, 226, 239}             # devenus archi, W-09, MCD, UC, séquence
# Paragraphes réécrits à la relecture (contradictions internes, renvois d'une
# section vers elle-même) : présents dans le dossier, sous une autre formulation.
REECRITS = {47, 49, 156, 261, 280, 457, 575, 585, 610}
EXCLUS = set(range(0, 22)) | TITRES_TRAME | COUPES | ASCII_REMPLACES | REECRITS

textes = [p.text for p in out.paragraphs if p.text.strip()]
for t in out.tables:
    for r in t.rows:
        for c in r.cells:
            textes += [p.text for p in c.paragraphs if p.text.strip()]
CLES = {cle(t) for t in textes}
BLOB = "\n".join(cle(t) for t in textes)

exacts = renvois = 0
manquants = []
for i, p in enumerate(src.paragraphs):
    if i in EXCLUS or not p.text.strip():
        continue
    if (p.style.name or "").startswith(("Heading", "Title")):
        continue                      # repris par la trame ou passé en gras
    if p.style.name == "Source Code":
        continue                      # rendu en image, contrôlé plus bas
    k = cle(p.text)
    if k in CLES:
        exacts += 1
    elif [m for m in re.split(r"§\s*[\d\.]+", p.text) if len(cle(m)) > 35] and all(
            cle(m) in BLOB for m in re.split(r"§\s*[\d\.]+", p.text) if len(cle(m)) > 35):
        renvois += 1                  # présent, renvoi réécrit
    else:
        manquants.append((i, p.style.name, p.text.strip()[:100]))

print(f"paragraphes de contenu : {exacts + renvois + len(manquants)}"
      f"  (+ {len(REECRITS)} réécrits à la relecture)")
print(f"  retrouvés à l'identique      : {exacts}")
print(f"  retrouvés, renvoi réécrit    : {renvois}")
print(f"  ABSENTS                      : {len(manquants)}")
for i, st, t in manquants:
    print(f"    [{i}] ({st}) {t}")

def empreinte(t):
    # les renvois sont neutralisés : « § 5.2.2 » devient « § 5.2 » dans le
    # dossier, ce qui ne change rien à la présence du tableau.
    return cle(" ".join(c.text for r in t.rows for c in r.cells), True)
emp_out = [empreinte(t) for t in out.tables]
# Le tableau 43 (§ 8) porte une cellule reformulée à la relecture : « durée 7
# jours » y est devenu « session de 7 jours côté serveur », le cookie portant
# lui un max_age de 30 jours.
TABLEAUX_RETOUCHES = {43}
tab_manquants = []
for i, t in enumerate(src.tables):
    if i in TABLEAUX_RETOUCHES:
        continue
    e = empreinte(t)
    if not any(e == o or e in o or o in e for o in emp_out):
        tab_manquants.append((i, " | ".join(c.text.strip()[:18] for c in t.rows[0].cells)))
print(f"\ntableaux : {len(src.tables)} source / {len(out.tables)} dossier, "
      f"{len(tab_manquants)} absents")
for i, h in tab_manquants:
    print(f"    [{i}] {h}")

# ── blocs de code rendus en images ────────────────────────────────────────
import hashlib, json, os
from code_nettoyage import nettoyer
IDX = "/home/user/HelpMeDraft---projet-fil-rouge/docs/dossier/fig/code/index.json"
manifeste = json.load(open(IDX))["blocs"] if os.path.exists(IDX) else {}
codes = [p.text for i, p in enumerate(src.paragraphs)
         if p.style.name == "Source Code" and i not in EXCLUS]
sans_image = [c.split("\n")[0][:60] for c in codes
              if hashlib.md5(nettoyer(c).encode()).hexdigest()[:12] not in manifeste]
print(f"\nblocs de code : {len(codes)} dans la source, "
      f"{sum(len(v) for v in manifeste.values())} images, {len(sans_image)} sans image")
for c in sans_image:
    print("    ", c)

print(f"\nimages insérées : {len(out.element.body.findall('.//' + qn('a:blip')))}")

reliquats = [p.text.strip() for p in out.paragraphs
             if any(m in p.text for m in ("[Nom de la société]", "Nom Prenom",
                                          "[Date]", "Trame du dossier projet",
                                          "Cliquez ou appuyez ici"))]
print(f"reliquats de la trame : {len(reliquats)} {reliquats}")

# ─────────────────── structure attendue par le référentiel ────────────────
# Les numéros des titres viennent de la numérotation automatique de la trame :
# ils ne figurent pas dans le texte. On compare donc les intitulés à ceux de la
# trame vierge, dans l'ordre, et on vérifie que chacun est suivi de contenu.
print("\nstructure")
trame = docx.Document("/home/user/HelpMeDraft---projet-fil-rouge/docs/dossier/"
                      "trame-dossier-projet-cda.docx")
# Sept des douze parties portent dans la trame le style « List Paragraph » ;
# elles sont dans la même liste numérotée que les autres. On les retient aussi.
attendus = [cle(p.text) for p in trame.paragraphs
            if p.text.strip() and ((p.style.name or "").startswith("Heading")
                                   or p._p.find(".//" + qn("w:numPr")) is not None)]
paras = out.paragraphs
idx = [i for i, p in enumerate(paras)
       if (p.style.name or "").startswith("Heading") and p.text.strip()]
obtenus = [cle(paras[i].text) for i in idx]
print(f"  titres : {len(obtenus)} / {len(attendus)} attendus — "
      f"{'conformes et dans l-ordre' if obtenus == attendus else 'ÉCART'}")
if obtenus != attendus:
    for a, b in zip(attendus, obtenus + [""] * len(attendus)):
        if a != b:
            print(f"    attendu {a[:50]} | obtenu {b[:50]}")
creux = []
for n, i in enumerate(idx):
    fin = idx[n + 1] if n + 1 < len(idx) else len(paras)
    contenu = [p for p in paras[i + 1:fin] if p.text.strip()]
    images = sum(1 for p in paras[i + 1:fin]
                 if p._p.findall(".//" + qn("a:blip")))
    if not contenu and not images:
        creux.append(paras[i].text.strip()[:60])
    print(f"    {paras[i].text.strip()[:58]:58} {len(contenu):3} paragraphes")
print(f"  parties sans contenu : {len(creux)} {creux}")
